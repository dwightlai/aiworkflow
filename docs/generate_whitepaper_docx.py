#!/usr/bin/env python3
"""Convert aiworkflow-whitepaper.md to DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "aiworkflow-whitepaper.md"
OUT_PATH = ROOT / "aiworkflow-whitepaper.docx"


def set_default_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_rich_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*[^*]+\*\*)", text.strip())
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if not line.strip().startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(set(cell) <= {"-", ":"} for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            value = row[c_idx] if c_idx < len(row) else ""
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = value
            if r_idx == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def convert() -> None:
    md = MD_PATH.read_text(encoding="utf-8")
    doc = Document()
    set_default_font(doc)

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    lines = md.splitlines()
    i = 0
    bullet_buffer: list[str] = []

    def flush_bullets() -> None:
        nonlocal bullet_buffer
        for item in bullet_buffer:
            add_rich_paragraph(doc, item, style="List Bullet")
        bullet_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped == "---":
            flush_bullets()
            i += 1
            continue

        if stripped.startswith("```"):
            flush_bullets()
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue

        if stripped.startswith("#"):
            flush_bullets()
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            doc.add_heading(title, level=min(level, 3))
            i += 1
            continue

        if stripped.startswith("!["):
            flush_bullets()
            match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
            if match:
                caption, rel_path = match.groups()
                image_path = (ROOT / rel_path).resolve()
                if image_path.exists():
                    doc.add_paragraph(caption, style="Caption")
                    doc.add_picture(str(image_path), width=Inches(5.8))
                    last = doc.paragraphs[-1]
                    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    add_rich_paragraph(doc, f"[图片缺失: {rel_path}]")
            i += 1
            continue

        if stripped.startswith("|"):
            flush_bullets()
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, parse_table(table_lines))
            continue

        if stripped.startswith("- "):
            bullet_buffer.append(stripped[2:].strip())
            i += 1
            continue

        if not stripped:
            flush_bullets()
            i += 1
            continue

        flush_bullets()
        add_rich_paragraph(doc, stripped)
        i += 1

    flush_bullets()
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    convert()
