from __future__ import annotations

import math
import re
import textwrap
from collections import defaultdict, deque
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from PIL import ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "aiworkflow-whitepaper.md"
OUT_DIR = ROOT / "docs" / "dist"
ASSET_DIR = OUT_DIR / "whitepaper-assets"
DOCX_OUT = OUT_DIR / "AI-Workflow-企业AI应用平台白皮书-v2.0.docx"
PDF_OUT = OUT_DIR / "AI-Workflow-企业AI应用平台白皮书-v2.0.pdf"
DOCX_FIXED_OUT = OUT_DIR / "AI-Workflow-企业AI应用平台白皮书-v2.0-修正版.docx"
PDF_FIXED_OUT = OUT_DIR / "AI-Workflow-企业AI应用平台白皮书-v2.0-修正版.pdf"
DOCX_RESULT = DOCX_OUT
PDF_RESULT = PDF_OUT

DIAGRAM_TITLES = [
    "AI Workflow 核心定位",
    "企业 AI 落地痛点与能力映射",
    "典型 AI 工作流编排",
    "企业知识库处理链路",
    "Chat 端智能体调用链路",
    "智能编研处理流程",
    "平台开放集成架构",
    "Jar/模块嵌入集成模式",
    "标准独立服务集成模式",
]


def clean_inline(text: str) -> str:
    text = text.strip()
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("<br>", "\n")
    return text


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_docx_font(document: Document) -> None:
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    document.styles["Normal"].font.size = Pt(10.5)


def add_docx_image(document: Document, image_path: Path, caption: str) -> None:
    if not image_path.exists():
        document.add_paragraph(f"[图片缺失：{image_path}]")
        return
    with PILImage.open(image_path) as img:
        width_px, height_px = img.size
    max_width = Inches(6.3)
    max_height = Inches(4.25)
    width = max_width
    if height_px and width_px:
        height = max_width * height_px / width_px
        if height > max_height:
            width = max_height * width_px / height_px
    document.add_picture(str(image_path), width=width)
    p = document.add_paragraph(caption)
    p.alignment = 1
    for run in p.runs:
        set_run_font(run, 9)
        run.font.color.rgb = RGBColor(100, 116, 139)


def add_docx_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_idx, row in enumerate(rows):
        cells = table.rows[row_idx].cells
        for col_idx, text in enumerate(row):
            cells[col_idx].text = clean_inline(text)
            for paragraph in cells[col_idx].paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, 9, row_idx == 0)
            if row_idx == 0:
                set_cell_shading(cells[col_idx], "EAF2FF")
    document.add_paragraph()


def parse_markdown_lines() -> list[str]:
    return SOURCE.read_text(encoding="utf-8").splitlines()


def font_path() -> str:
    candidates = [
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    return candidates[-1]


def text_bbox(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont) -> tuple[int, int]:
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=3)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_label(label: str, max_chars: int = 13) -> str:
    if label == "AI Workflow":
        return label
    if label == "OpenAPI + Java Client":
        return "OpenAPI +\nJava Client"
    if label == "运行记录 + 节点轨迹":
        return "运行记录 +\n节点轨迹"
    if len(label) <= max_chars:
        return label
    pieces = textwrap.wrap(label, width=max_chars, break_long_words=False, replace_whitespace=False)
    return "\n".join(pieces) if pieces else label


def parse_mermaid(code: str) -> tuple[str, dict[str, str], list[tuple[str, str, str]]]:
    direction = "LR"
    labels: dict[str, str] = {}
    edges: list[tuple[str, str, str]] = []

    node_pattern = re.compile(r'([A-Za-z0-9_]+)\s*\[\s*"([^"]+)"\s*\]')
    for raw in code.splitlines():
        line = raw.strip().rstrip(";")
        if not line or line.startswith("%%"):
            continue
        if line.startswith("flowchart") or line.startswith("graph"):
            parts = line.split()
            if len(parts) > 1:
                direction = parts[1].strip().upper()
            continue

        for node_id, label in node_pattern.findall(line):
            labels[node_id] = label

        edge_match = re.search(r"([A-Za-z0-9_]+)(?:\s*\[\s*\"([^\"]+)\"\s*\])?\s*-->\s*(?:\|([^|]+)\|\s*)?([A-Za-z0-9_]+)(?:\s*\[\s*\"([^\"]+)\"\s*\])?", line)
        if edge_match:
            source, source_label, edge_label, target, target_label = edge_match.groups()
            if source_label:
                labels[source] = source_label
            if target_label:
                labels[target] = target_label
            labels.setdefault(source, source)
            labels.setdefault(target, target)
            edges.append((source, target, edge_label or ""))

    return direction, labels, edges


def compute_levels(nodes: list[str], edges: list[tuple[str, str, str]]) -> dict[str, int]:
    indegree = {node: 0 for node in nodes}
    children: dict[str, list[str]] = defaultdict(list)
    feedback_target = edges[0][0] if edges else ""
    for source, target, _ in edges:
        if target == feedback_target:
            continue
        children[source].append(target)
        indegree[target] = indegree.get(target, 0) + 1
        indegree.setdefault(source, 0)

    queue = deque([node for node in nodes if indegree.get(node, 0) == 0])
    levels = {node: 0 for node in nodes}
    while queue:
        node = queue.popleft()
        for child in children[node]:
            levels[child] = max(levels.get(child, 0), levels[node] + 1)
            indegree[child] -= 1
            if indegree[child] == 0:
                queue.append(child)
    if all(level == 0 for level in levels.values()) and edges:
        # Fall back to first-edge source when a diagram contains a feedback loop.
        levels[edges[0][0]] = 0
        for _ in range(len(nodes)):
            changed = False
            for source, target, _ in edges:
                if target == edges[0][0]:
                    continue
                new_level = levels.get(source, 0) + 1
                if new_level > levels.get(target, 0):
                    levels[target] = new_level
                    changed = True
            if not changed:
                break
    return levels


def arrow(draw: ImageDraw.ImageDraw, start: tuple[float, float], end: tuple[float, float], color: str = "#333333") -> None:
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=3)
    angle = math.atan2(y2 - y1, x2 - x1)
    size = 13
    points = [
        (x2, y2),
        (x2 - size * math.cos(angle - math.pi / 6), y2 - size * math.sin(angle - math.pi / 6)),
        (x2 - size * math.cos(angle + math.pi / 6), y2 - size * math.sin(angle + math.pi / 6)),
    ]
    draw.polygon(points, fill=color)


def render_mermaid_diagram(code: str, index: int) -> Path:
    direction, labels, edges = parse_mermaid(code)
    nodes = list(labels.keys())
    levels = compute_levels(nodes, edges)
    by_level: dict[int, list[str]] = defaultdict(list)
    for node in nodes:
        by_level[levels.get(node, 0)].append(node)
    level_count = max(by_level.keys(), default=0) + 1

    margin_x = 90
    margin_y = 55
    box_w = 210
    box_h = 78
    max_nodes_in_level = max((len(v) for v in by_level.values()), default=1)
    if direction == "TB":
        width = max(1100, 2 * margin_x + max_nodes_in_level * (box_w + 24))
        height = max(360, 130 + level_count * 104, 130 + max_nodes_in_level * 90)
    else:
        width = max(1100, 2 * margin_x + level_count * (box_w + 32))
        height = max(360, 130 + max_nodes_in_level * 100)

    image = PILImage.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(font_path(), 21)
    edge_font = ImageFont.truetype(font_path(), 15)

    positions: dict[str, tuple[float, float]] = {}
    if direction == "TB":
        for level, level_nodes in by_level.items():
            y = margin_y + level * ((height - 2 * margin_y) / max(level_count - 1, 1))
            gap = (width - 2 * margin_x) / max(len(level_nodes), 1)
            for idx, node in enumerate(level_nodes):
                positions[node] = (margin_x + gap * idx + gap / 2, y)
    else:
        for level, level_nodes in by_level.items():
            x = margin_x + level * ((width - 2 * margin_x) / max(level_count - 1, 1))
            gap = (height - 2 * margin_y) / max(len(level_nodes), 1)
            for idx, node in enumerate(level_nodes):
                positions[node] = (x, margin_y + gap * idx + gap / 2)

    for source, target, edge_label in edges:
        sx, sy = positions[source]
        tx, ty = positions[target]
        if direction == "TB":
            start = (sx, sy + box_h / 2)
            end = (tx, ty - box_h / 2)
        else:
            start = (sx + box_w / 2, sy)
            end = (tx - box_w / 2, ty)
        arrow(draw, start, end)
        if edge_label:
            mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
            tw, th = text_bbox(draw, edge_label, edge_font)
            draw.rounded_rectangle((mx - tw / 2 - 6, my - th / 2 - 4, mx + tw / 2 + 6, my + th / 2 + 4), radius=4, fill="white", outline="#CBD5E1")
            draw.text((mx - tw / 2, my - th / 2), edge_label, font=edge_font, fill="#334155")

    for node, (cx, cy) in positions.items():
        label = wrap_label(labels[node])
        x1, y1 = cx - box_w / 2, cy - box_h / 2
        x2, y2 = cx + box_w / 2, cy + box_h / 2
        draw.rounded_rectangle((x1, y1, x2, y2), radius=2, fill="#F5F0FF", outline="#8B5CF6", width=2)
        tw, th = text_bbox(draw, label, font)
        draw.multiline_text((cx - tw / 2, cy - th / 2), label, font=font, fill="#1E1B4B", align="center", spacing=3)

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    output = ASSET_DIR / f"diagram-{index:02d}.png"
    image.save(output)
    return output


def export_docx(lines: list[str]) -> None:
    global DOCX_RESULT
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    set_docx_font(document)

    title = clean_inline(lines[0].lstrip("# "))
    document.add_heading(title, level=0)
    subtitle = document.add_paragraph("企业 AI 应用平台白皮书 · v2.0 · 2026-06-17")
    subtitle.runs[0].font.color.rgb = RGBColor(71, 85, 105)
    document.add_paragraph("面向业务负责人、信息化负责人、架构师、实施团队")
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    i = 1
    in_code = False
    code_lang = ""
    code_lines: list[str] = []
    diagram_index = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = stripped[3:].strip()
                code_lines = []
            else:
                if code_lang == "mermaid":
                    diagram_index += 1
                    diagram_path = render_mermaid_diagram("\n".join(code_lines), diagram_index)
                    caption = f"图 {diagram_index}：{DIAGRAM_TITLES[diagram_index - 1] if diagram_index <= len(DIAGRAM_TITLES) else '流程图'}"
                    add_docx_image(document, diagram_path, caption)
                else:
                    p = document.add_paragraph()
                    run = p.add_run("\n".join(code_lines))
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                in_code = False
                code_lang = ""
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped or stripped == "---":
            i += 1
            continue

        image_match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            caption = image_match.group(1)
            image_path = (SOURCE.parent / image_match.group(2)).resolve()
            add_docx_image(document, image_path, caption)
            i += 1
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                current = lines[i].strip()
                table_lines.append(current)
                i += 1
            rows = []
            for table_line in table_lines:
                row = [cell.strip() for cell in table_line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row):
                    rows.append(row)
            add_docx_table(document, rows)
            continue

        if stripped.startswith("#"):
            level = min(stripped.count("#"), 3)
            text = clean_inline(stripped.lstrip("# "))
            document.add_heading(text, level=level)
        elif stripped.startswith("- "):
            document.add_paragraph(clean_inline(stripped[2:]), style="List Bullet")
        elif re.match(r"^\d+\.\s+", stripped):
            document.add_paragraph(clean_inline(re.sub(r"^\d+\.\s+", "", stripped)), style="List Number")
        elif stripped.startswith(">"):
            document.add_paragraph(clean_inline(stripped.lstrip("> ")))
        else:
            document.add_paragraph(clean_inline(stripped))
        i += 1

    try:
        document.save(DOCX_OUT)
        DOCX_RESULT = DOCX_OUT
    except PermissionError:
        document.save(DOCX_FIXED_OUT)
        DOCX_RESULT = DOCX_FIXED_OUT


def setup_pdf_styles():
    pdfmetrics.registerFont(TTFont("SimHei", str(font_path())))
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("title", parent=styles["Title"], fontName="SimHei", fontSize=22, leading=30, alignment=TA_CENTER, spaceAfter=18),
        "h1": ParagraphStyle("h1", parent=styles["Heading1"], fontName="SimHei", fontSize=17, leading=23, textColor=colors.HexColor("#0f172a"), spaceBefore=12, spaceAfter=8),
        "h2": ParagraphStyle("h2", parent=styles["Heading2"], fontName="SimHei", fontSize=14, leading=20, textColor=colors.HexColor("#1d4ed8"), spaceBefore=10, spaceAfter=6),
        "h3": ParagraphStyle("h3", parent=styles["Heading3"], fontName="SimHei", fontSize=12, leading=17, textColor=colors.HexColor("#334155"), spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("body", parent=styles["BodyText"], fontName="SimHei", fontSize=9.6, leading=15, alignment=TA_LEFT, spaceAfter=5),
        "caption": ParagraphStyle("caption", parent=styles["BodyText"], fontName="SimHei", fontSize=8.5, leading=12, alignment=TA_CENTER, textColor=colors.HexColor("#64748b"), spaceAfter=8),
        "code": ParagraphStyle("code", parent=styles["Code"], fontName="Courier", fontSize=7.2, leading=9, backColor=colors.HexColor("#f8fafc"), borderColor=colors.HexColor("#e2e8f0"), borderWidth=0.5, borderPadding=5, spaceAfter=8),
    }


def pdf_escape(text: str) -> str:
    return clean_inline(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def add_pdf_image(story, image_path: Path, caption: str, styles) -> None:
    if not image_path.exists():
        story.append(Paragraph(f"图片缺失：{image_path}", styles["body"]))
        return
    max_w = 17.2 * cm
    max_h = 10.5 * cm
    with PILImage.open(image_path) as img:
        w, h = img.size
    scale = min(max_w / w, max_h / h, 1)
    story.append(Image(str(image_path), width=w * scale, height=h * scale))
    story.append(Paragraph(caption, styles["caption"]))


def add_pdf_table(story, rows: list[list[str]], styles) -> None:
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    normalized = [row + [""] * (max_cols - len(row)) for row in rows]
    data = [[Paragraph(pdf_escape(cell), styles["body"]) for cell in row] for row in normalized]
    col_width = (17.2 * cm) / max_cols
    table = Table(data, colWidths=[col_width] * max_cols, hAlign="LEFT", repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf2ff")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    story.append(Spacer(1, 8))


def export_pdf(lines: list[str]) -> None:
    global PDF_RESULT
    styles = setup_pdf_styles()
    doc = SimpleDocTemplate(str(PDF_OUT), pagesize=A4, rightMargin=1.8 * cm, leftMargin=1.8 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm)
    story = [
        Paragraph("AI Workflow 企业 AI 应用平台白皮书", styles["title"]),
        Paragraph("v2.0 · 2026-06-17 · 面向业务负责人、信息化负责人、架构师、实施团队", styles["caption"]),
        Spacer(1, 12),
    ]

    i = 1
    in_code = False
    code_lang = ""
    code_lines: list[str] = []
    diagram_index = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = stripped[3:].strip()
                code_lines = []
            else:
                if code_lang == "mermaid":
                    diagram_index += 1
                    diagram_path = render_mermaid_diagram("\n".join(code_lines), diagram_index)
                    caption = f"图 {diagram_index}：{DIAGRAM_TITLES[diagram_index - 1] if diagram_index <= len(DIAGRAM_TITLES) else '流程图'}"
                    add_pdf_image(story, diagram_path, caption, styles)
                else:
                    story.append(Paragraph(pdf_escape("\n".join(code_lines)).replace("\n", "<br/>"), styles["code"]))
                in_code = False
                code_lang = ""
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped or stripped == "---":
            i += 1
            continue
        image_match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            add_pdf_image(story, (SOURCE.parent / image_match.group(2)).resolve(), image_match.group(1), styles)
            i += 1
            continue
        if stripped.startswith("|") and "|" in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for table_line in table_lines:
                row = [cell.strip() for cell in table_line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row):
                    rows.append(row)
            add_pdf_table(story, rows, styles)
            continue
        if stripped.startswith("#"):
            level = stripped.count("#")
            text = pdf_escape(stripped.lstrip("# "))
            story.append(Paragraph(text, styles["h1" if level == 1 else "h2" if level == 2 else "h3"]))
        elif stripped.startswith("- "):
            story.append(ListFlowable([ListItem(Paragraph(pdf_escape(stripped[2:]), styles["body"]))], bulletType="bullet", start="circle"))
        elif re.match(r"^\d+\.\s+", stripped):
            story.append(Paragraph("• " + pdf_escape(re.sub(r"^\d+\.\s+", "", stripped)), styles["body"]))
        elif stripped.startswith(">"):
            story.append(Paragraph(pdf_escape(stripped.lstrip("> ")), styles["body"]))
        else:
            story.append(Paragraph(pdf_escape(stripped), styles["body"]))
        i += 1
    try:
        doc.build(story)
        PDF_RESULT = PDF_OUT
    except PermissionError:
        fallback_doc = SimpleDocTemplate(str(PDF_FIXED_OUT), pagesize=A4, rightMargin=1.8 * cm, leftMargin=1.8 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm)
        fallback_doc.build(story)
        PDF_RESULT = PDF_FIXED_OUT


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    lines = parse_markdown_lines()
    export_docx(lines)
    export_pdf(lines)
    print(DOCX_RESULT)
    print(PDF_RESULT)


if __name__ == "__main__":
    main()
